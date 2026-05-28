import os
import logging
import tempfile
from datetime import datetime

from openpyxl import Workbook

from telegram.ext import ContextTypes
from config import GROUP_CHAT_ID, ADMIN_ID, CHECKER_TIMEOUT_SEC

logger = logging.getLogger(__name__)


async def is_topic_valid(context: ContextTypes.DEFAULT_TYPE, topic_id: int | None) -> bool:
    # Trust the DB value — if the topic is gone, the next real send will raise
    # BadRequest("message thread not found") which xodim_chat_handler already handles.
    # Sending a probe message here clutters the admin group on every /start.
    return bool(topic_id)


async def check_sla_timeout(context: ContextTypes.DEFAULT_TYPE):
    from database import get_xabar, get_group_info
    job      = context.job
    x_ism    = job.data["x_ism"]
    reminder = job.data.get("reminder", 1)

    group_id = job.data.get("group_id")
    task_id  = job.data.get("task_id")  # eski xabarlar uchun

    if group_id:
        row = await get_group_info(group_id)
        if not row or row[4] != "kutilmoqda":
            return
    elif task_id:
        row = await get_xabar(task_id)
        if not row or row[3] != "kutilmoqda":
            return
    else:
        return

    minutes    = 15 * reminder
    ref_id     = group_id or task_id
    ogohlantirish = (
        f"🚨 *DIQQAT! #SLA Nazorati ({reminder}-eslatma)*\n\n"
        f"⚠️ #{ref_id}-sonli topshiriq kelganiga *{minutes} daqiqa* bo'ldi, "
        f"biroq haligacha belgilanmadi!\n"
        f"👤 Xodim: {x_ism}"
    )
    try:
        await context.bot.send_message(
            chat_id=ADMIN_ID, text=ogohlantirish, parse_mode="Markdown"
        )
    except Exception as e:
        logger.error(f"Admin SLA xabari yuborishda xato: {e}")


async def daily_report_job(context: ContextTypes.DEFAULT_TYPE):
    # Dushanbada haftalik hisobot yuboriladi, kunlik shart emas
    if datetime.now().weekday() == 0:
        return
    from database import get_statistika, get_kunlik_statistika
    stat   = await get_statistika()
    kunlik = await get_kunlik_statistika()
    today  = datetime.now().strftime("%d.%m.%Y")

    matn = (
        f"📅 *{today} — Kunlik Hisobot*\n\n"
        f"📥 Jami xabarlar: *{stat['jami']}*\n"
        f"✅ Bajarilgan: *{stat['bajarilgan']}*\n"
        f"⏳ Kutilmoqda: *{stat['kutilmoqda']}*\n"
        f"⏱ O'rtacha vaqt: *{stat['ortacha']} daqiqa*"
    )
    if kunlik:
        matn += "\n\n*📊 Bugungi faollik:*\n"
        for row in kunlik:
            matn += f"• {row[0]} ({row[1]}): {row[2]} ta, {row[3]} bajarildi\n"
    else:
        matn += "\n\n_Bugun hech qanday faollik yo'q._"

    try:
        await context.bot.send_message(
            chat_id=ADMIN_ID, text=matn, parse_mode="Markdown"
        )
    except Exception as e:
        logger.error(f"Kunlik hisobot yuborishda xato: {e}")


async def urgency_timeout_job(context: ContextTypes.DEFAULT_TYPE):
    """60 soniyada urgency tanlanmasa oddiy deb belgilab checker ga xabar yuboradi."""
    from database import set_urgency, get_latest_group_fwd_id, get_group_info
    data       = context.job.data
    group_id   = data["group_id"]
    checker_id = data["checker_id"]
    ism        = data["ism"]

    await set_urgency(group_id, "oddiy")
    try:
        fwd_id = await get_latest_group_fwd_id(group_id)
        if fwd_id:
            await context.bot.forward_message(
                chat_id=checker_id, from_chat_id=GROUP_CHAT_ID, message_id=fwd_id
            )
        from keyboards import checker_sorov_kb
        await context.bot.send_message(
            chat_id=checker_id,
            text=(
                f"📋 *{ism}* (Agent) #{group_id} topshiriq yubordi.\n"
                f"🟢 _Urgency: Oddiy_\nKo'rib chiqing:"
            ),
            parse_mode="Markdown",
            reply_markup=checker_sorov_kb(group_id),
        )
    except Exception as e:
        logger.warning(f"Urgency timeout checker xabari xato: {e}")


async def checker_timeout_job(context: ContextTypes.DEFAULT_TYPE):
    """Checker 30 daqiqada javob bermasa adminga xabar yuboradi."""
    from database import get_group_info
    data     = context.job.data
    group_id = data["group_id"]
    ism      = data["ism"]

    group = await get_group_info(group_id)
    if not group or group[4] != "kutilmoqda":
        return  # Allaqachon bajarilgan

    try:
        await context.bot.send_message(
            chat_id=ADMIN_ID,
            text=(
                f"⏰ *Checker javob bermadi!*\n\n"
                f"#{group_id}-sonli topshiriq uchun checker 30 daqiqada javob bermadi.\n"
                f"👤 Agent: *{ism}*"
            ),
            parse_mode="Markdown",
        )
    except Exception as e:
        logger.warning(f"Checker timeout admin xabari xato: {e}")


async def weekly_report_job(context: ContextTypes.DEFAULT_TYPE):
    """Dushanba 09:00 — haftalik hisobot (kunlik hisobotni almashtiradi dushanbada)."""
    from database import get_checker_weekly_stats, get_statistika
    rows  = await get_checker_weekly_stats()
    stat  = await get_statistika()
    today = datetime.now().strftime("%d.%m.%Y")

    matn = (
        f"📅 *Haftalik Hisobot ({today})*\n\n"
        f"📊 Jami: *{stat['jami']}* | ✅ *{stat['bajarilgan']}* | ⏳ *{stat['kutilmoqda']}*\n"
        f"⏱ O'rtacha: *{stat['ortacha']} daqiqa*\n\n"
        f"━━━━━━━━━━━━━━\n\n"
    )
    if not rows:
        matn += "_Bu hafta biriktirish yo'q yoki faollik kuzatilmadi._"
    else:
        for r in rows:
            checker_ism, agents, topshiriq, bajarildi, avg_r = r
            stars = ("⭐" * round(avg_r)) if avg_r else "—"
            matn += (
                f"👤 *{checker_ism}*\n"
                f"  Agentlar: {agents} ta  |  Topshiriqlar: {topshiriq}\n"
                f"  Bajarildi: {bajarildi or 0}  |  Reyting: {stars}\n\n"
            )
    try:
        await context.bot.send_message(chat_id=ADMIN_ID, text=matn, parse_mode="Markdown")
    except Exception as e:
        logger.error(f"Haftalik hisobot yuborishda xato: {e}")


async def agent_reminder_job(context: ContextTypes.DEFAULT_TYPE):
    """Har kuni 17:00 — bugun xabar yubormagan agentlarning checker larini xabardor qilish."""
    from database import get_agents_without_messages_today, get_biriktirish
    agents = await get_agents_without_messages_today()
    for agent_id, agent_ism in agents:
        checker_id = await get_biriktirish(agent_id)
        if not checker_id:
            continue
        try:
            await context.bot.send_message(
                chat_id=checker_id,
                text=(
                    f"⏰ *Eslatma:* *{agent_ism}* bugun hech qanday topshiriq yubormaganlar.\n"
                    f"Tekshirib ko'ring."
                ),
                parse_mode="Markdown",
            )
        except Exception as e:
            logger.warning(f"Agent reminder xato (checker={checker_id}): {e}")


async def generate_excel(x_data: list, m_data: list) -> str:
    suffix = f"_Dusel_Hisobot_{datetime.now().strftime('%d_%m_%Y')}.xlsx"
    fd, filename = tempfile.mkstemp(suffix=suffix)
    os.close(fd)

    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Xodimlar"
    ws1.append([
        "User ID", "Ism", "Lavozim", "Kod", "Filial",
        "Tel 1", "Tel 2", "Tug'ilgan kun", "Status", "Ro'yxatdan o'tgan sana",
    ])
    for row in x_data:
        ws1.append(list(row))

    ws2 = wb.create_sheet("Topshiriqlar_SLA")
    ws2.append([
        "Vazifa ID", "Xodim ID", "Xodim Ismi", "Filial",
        "Xabar Turi", "Xabar Kelgan Vaqt", "Xabar Holati", "Bajarilgan Vaqt",
    ])
    for row in m_data:
        ws2.append(list(row))

    wb.save(filename)
    return filename


async def generate_full_excel(data: dict) -> str:
    suffix = f"_DuselDB_{datetime.now().strftime('%d_%m_%Y')}.xlsx"
    fd, filename = tempfile.mkstemp(suffix=suffix)
    os.close(fd)

    wb = Workbook()

    # Sheet 1: Xodimlar
    ws = wb.active
    ws.title = "Xodimlar"
    ws.append(["User ID", "Ism", "Lavozim", "Kod", "Filial",
               "Tel 1", "Tel 2", "Tug'ilgan kun", "Status", "Sana"])
    for row in data.get("xodimlar", []):
        ws.append(list(row))

    # Sheet 2: Topshiriqlar
    ws2 = wb.create_sheet("Topshiriqlar")
    ws2.append(["ID", "User ID", "Ism", "Filial", "Topic ID",
                "Holat", "Urgency", "Kelgan vaqt", "Javob vaqt"])
    for row in data.get("topshiriqlar", []):
        ws2.append(list(row))

    # Sheet 3: Klientlar
    ws3 = wb.create_sheet("Klientlar")
    ws3.append(["ID", "Firma nomi", "Tel 1", "Tel 2", "INN", "Orienter",
                "Lat", "Lon", "Manzil", "Kategoriya", "Do'kon turi",
                "Distributor", "Agent kodi", "Vizit kun", "Chastota",
                "Limit", "Brendlar", "Holat", "Rad sababi", "Sana", "Supervisor ID"])
    for row in data.get("klientlar", []):
        ws3.append(list(row))

    # Sheet 4: So'rovlar
    ws4 = wb.create_sheet("Sorovlar")
    ws4.append(["ID", "Agent ID", "Agent Ism", "Tur", "Do'kon nomi",
                "Yangi qiymat", "Lat", "Lon", "Izoh", "Status", "Sana"])
    for row in data.get("sorovlar", []):
        ws4.append(list(row))

    # Sheet 5: Biriktirish
    ws5 = wb.create_sheet("Biriktirish")
    ws5.append(["Agent ID", "Agent Ism", "Checker ID", "Checker Ism",
                "Filial", "Status"])
    for row in data.get("biriktirish", []):
        ws5.append(list(row))

    # Sheet 6: Baholash
    ws6 = wb.create_sheet("Baholash")
    ws6.append(["ID", "Guruh ID", "Checker ID", "Checker Ism",
                "Agent ID", "Agent Ism", "Yulduz", "Vaqt"])
    for row in data.get("baholash", []):
        ws6.append(list(row))

    # Sheet 7: Audit Log
    ws7 = wb.create_sheet("Audit Log")
    ws7.append(["ID", "User ID", "Ism", "Rol", "Amal", "Nishon",
                "Eski qiymat", "Yangi qiymat", "Status", "Vaqt"])
    for row in data.get("audit_log", []):
        ws7.append(list(row))

    wb.save(filename)
    return filename


async def pending_sorovlar_alert_job(context: ContextTypes.DEFAULT_TYPE):
    from database import get_pending_sorovlar_by_supervisor
    sup_counts = await get_pending_sorovlar_by_supervisor()
    if not sup_counts:
        return
    for sup_id, count in sup_counts.items():
        try:
            await context.bot.send_message(
                chat_id=sup_id,
                text=(
                    f"⏰ *{count} ta so'rov tasdiqlanishingizni kutmoqda\\!*\n"
                    "_Agentlar tomonidan yuborilgan va hali tasdiqlanmagan so'rovlar\\._"
                ),
                parse_mode="MarkdownV2",
            )
        except Exception as e:
            logger.warning(f"Pending sorov alert xato (supervisor={sup_id}): {e}")


async def generate_klientlar_excel(rows: list) -> str:
    suffix = f"_Ochilgan_Klientlar_{datetime.now().strftime('%d_%m_%Y')}.xlsx"
    fd, filename = tempfile.mkstemp(suffix=suffix)
    os.close(fd)

    wb = Workbook()
    ws = wb.active
    ws.title = "Ochilgan Klientlar"
    ws.append([
        "ID", "Firma nomi", "Tel 1", "Tel 2", "INN",
        "Orienter", "Lat", "Lon", "Manzil",
        "Kategoriya", "Do'kon turi", "Distributor", "Agent kodi",
        "Limit", "Brendlar", "Holat", "Sana", "Supervisor ID",
    ])
    for row in rows:
        ws.append(list(row))

    wb.save(filename)
    return filename
